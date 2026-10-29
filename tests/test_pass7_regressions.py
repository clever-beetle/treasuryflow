import io
import uuid
import csv
from datetime import datetime, timezone, timedelta


def _register_and_login(client, prefix):
    username = f"{prefix}_{uuid.uuid4().hex[:6]}"
    password = "StrongPass123"
    client.post('/register', data={
        'fullname': 'Pass Seven',
        'username': username,
        'email': f"{username}@test.com",
        'password': password,
        'confirm_password': password
    }, follow_redirects=True)
    client.post('/login', data={'username': username, 'password': password}, follow_redirects=True)
    return username


def test_spreadsheet_safe_helper():
    from utils import spreadsheet_safe
    assert spreadsheet_safe('=HYPERLINK("http://evil","x")').startswith("'=")
    assert spreadsheet_safe('+62812') == "'+62812"
    assert spreadsheet_safe('-5') == "'-5"
    assert spreadsheet_safe('@SUM(A1)') == "'@SUM(A1)"
    assert spreadsheet_safe('Makan siang') == 'Makan siang'
    assert spreadsheet_safe(15000) == 15000
    assert spreadsheet_safe(None) is None


def test_wib_now_is_utc_plus_7_regardless_of_server_tz():
    from utils import wib_now
    expected = datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(hours=7)
    assert abs((wib_now() - expected).total_seconds()) < 5
    assert wib_now().tzinfo is None


def test_exports_neutralize_formula_injection(client):
    username = _register_and_login(client, 'pass7_export')
    from utils import get_db
    with client.application.app_context():
        db = get_db()
        user_id = db.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()['id']
        db.execute("INSERT INTO accounts (user_id, name, initial_balance, current_balance, account_type) VALUES (?, '[BANK] BCA', 1000000, 1000000, 'asset')", (user_id,))
        db.commit()
        acc_id = db.execute("SELECT id FROM accounts WHERE user_id = ?", (user_id,)).fetchone()['id']

    payload = '=HYPERLINK("http://evil.example","klik")'
    client.post('/transactions/add', data={
        'date': '2026-10-01', 'type': 'expense', 'amount': '10.000',
        'category': 'Belanja', 'account_id': acc_id, 'description': payload
    }, follow_redirects=True)

    res_csv = client.get('/transactions/export_csv')
    assert res_csv.status_code == 200
    rows = list(csv.reader(io.StringIO(res_csv.data.decode('utf-8-sig'))))
    descs = [r[3] for r in rows[1:]]
    assert "'" + payload in descs
    assert payload not in descs

    res_xlsx = client.get('/transactions/export_excel')
    assert res_xlsx.status_code == 200
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.data))
    ws = wb.active
    values = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
    assert "'" + payload in values
    # Must be stored as plain text, never as a formula cell
    for r in range(2, ws.max_row + 1):
        assert ws.cell(row=r, column=4).data_type != 'f'


def test_deleted_income_category_falls_back_to_valid_income_default(client):
    username = _register_and_login(client, 'pass7_cat')
    from utils import get_db, CATEGORIES
    with client.application.app_context():
        db = get_db()
        user_id = db.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()['id']
        db.execute("INSERT INTO accounts (user_id, name, initial_balance, current_balance, account_type) VALUES (?, '[BANK] BCA', 0, 0, 'asset')", (user_id,))
        db.commit()
        acc_id = db.execute("SELECT id FROM accounts WHERE user_id = ?", (user_id,)).fetchone()['id']

    client.post('/settings', data={'action': 'add_category', 'type': 'income', 'name': 'Freelance'}, follow_redirects=True)
    client.post('/transactions/add', data={
        'date': '2026-10-01', 'type': 'income', 'amount': '750.000',
        'category': 'Freelance', 'account_id': acc_id, 'description': 'Proyek web'
    }, follow_redirects=True)
    with client.application.app_context():
        db = get_db()
        cat_id = db.execute("SELECT id FROM user_categories WHERE user_id = ? AND name = 'Freelance'", (user_id,)).fetchone()['id']

    client.post('/settings', data={'action': 'delete_category', 'category_id': cat_id}, follow_redirects=True)
    with client.application.app_context():
        db = get_db()
        cat = db.execute("SELECT category FROM transactions WHERE user_id = ? AND description = 'Proyek web'", (user_id,)).fetchone()['category']
    assert cat == 'Lainnya (Pemasukan)'
    assert cat in CATEGORIES['income']
